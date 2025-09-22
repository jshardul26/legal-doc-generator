"""
Postprocessing Service for Legal Document Generator.

This module handles document formatting, validation, cleanup, and conversion
to various output formats after LLM generation.
"""

import re
import logging
import os
import tempfile
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, Tuple, Union
from dataclasses import dataclass
from enum import Enum
import hashlib
import base64
from pathlib import Path

# Document processing libraries
try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logger.warning("ReportLab not available - PDF generation disabled")

try:
    from docx import Document as DocxDocument
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False
    logger.warning("python-docx not available - DOCX generation disabled")

from app.models.document import DocumentFormat, DocumentType
from app.core.config import settings

# Configure logging
logger = logging.getLogger(__name__)


class PostprocessingError(Exception):
    """Base exception for postprocessing errors."""
    pass


class FormatNotSupportedError(PostprocessingError):
    """Exception for unsupported output formats."""
    pass


class ValidationFailedError(PostprocessingError):
    """Exception for document validation failures."""
    pass


@dataclass
class ProcessedDocument:
    """Result of document postprocessing."""
    content: str
    format: DocumentFormat
    file_path: Optional[str] = None
    file_size: Optional[int] = None
    metadata: Dict[str, Any] = None
    validation_results: Dict[str, Any] = None
    processing_time: float = 0.0
    warnings: List[str] = None


class DocumentCleaner:
    """Handles document content cleaning and normalization."""
    
    @staticmethod
    def clean_document_content(content: str) -> str:
        """Clean and normalize generated document content."""
        if not content:
            return ""
        
        # Remove excessive whitespace while preserving document structure
        content = re.sub(r'\n\s*\n\s*\n', '\n\n', content)  # Max 2 consecutive newlines
        content = re.sub(r'[ \t]+', ' ', content)  # Multiple spaces to single space
        content = re.sub(r'[ \t]*\n', '\n', content)  # Remove trailing spaces on lines
        
        # Fix common formatting issues
        content = re.sub(r'\n+([A-Z][A-Z\s]+:)', r'\n\n\1', content)  # Section headers
        content = re.sub(r'([.!?])\s*([A-Z])', r'\1 \2', content)  # Space after sentences
        content = re.sub(r'([a-z]),([A-Z])', r'\1, \2', content)  # Space after commas
        
        # Fix legal document formatting
        content = re.sub(r'\b(WHEREAS|NOW THEREFORE|IN WITNESS WHEREOF)\b', 
                        r'\n\n\1', content)  # Legal clauses on new lines
        content = re.sub(r'(\d+\.)\s*([A-Z])', r'\n\1 \2', content)  # Numbered sections
        
        # Clean up signature blocks
        content = re.sub(r'(Signature:?\s*_+)', r'\n\n\1', content)
        content = re.sub(r'(Date:?\s*_+)', r'\n\1', content)
        
        # Remove placeholder formatting artifacts
        content = re.sub(r'\[PLACEHOLDER[^\]]*\]', 
                        lambda m: m.group(0).replace('_', ' '), content)
        
        return content.strip()
    
    @staticmethod
    def add_legal_disclaimer(content: str, document_type: DocumentType) -> str:
        """Add appropriate legal disclaimer to document."""
        
        disclaimer_templates = {
            DocumentType.LEASE_AGREEMENT: """
LEGAL DISCLAIMER: This document is a template and may not comply with all local, 
state, or federal laws. It should be reviewed by a qualified attorney before use. 
The parties acknowledge that they have been advised to seek independent legal counsel.
""",
            DocumentType.NDA: """
LEGAL DISCLAIMER: This Non-Disclosure Agreement template should be reviewed by 
qualified legal counsel to ensure it meets your specific needs and complies with 
applicable laws in your jurisdiction.
""",
            DocumentType.SERVICE_AGREEMENT: """
LEGAL DISCLAIMER: This service agreement template is provided for informational 
purposes only and should be reviewed by qualified legal counsel before execution. 
Laws vary by jurisdiction and specific circumstances.
""",
            DocumentType.EMPLOYMENT_CONTRACT: """
LEGAL DISCLAIMER: This employment contract template must be reviewed by qualified 
legal counsel and HR professionals to ensure compliance with applicable employment 
laws, regulations, and local requirements.
""",
            "default": """
LEGAL DISCLAIMER: This document template is provided for informational purposes 
only and should be reviewed by qualified legal counsel before use. Laws and 
regulations vary by jurisdiction and may affect the validity and enforceability 
of this document.
"""
        }
        
        disclaimer = disclaimer_templates.get(document_type, disclaimer_templates["default"])
        
        return f"{content}\n\n{'-' * 80}\n{disclaimer.strip()}\n{'-' * 80}"
    
    @staticmethod
    def format_signature_blocks(content: str, parties: List[str] = None) -> str:
        """Add properly formatted signature blocks."""
        
        if not parties:
            # Try to extract party names from content
            parties = DocumentCleaner._extract_party_names(content)
        
        if not parties:
            parties = ["Party 1", "Party 2"]  # Default fallback
        
        signature_block = "\n\nSIGNATURES:\n\n"
        
        for party in parties:
            signature_block += f"""
{party}:

Signature: _________________________________  Date: ________________

Print Name: _________________________________

Title: _________________________________

"""
        
        # Add witness section if needed
        signature_block += """

WITNESS:

Signature: _________________________________  Date: ________________

Print Name: _________________________________

"""
        
        return content + signature_block
    
    @staticmethod
    def _extract_party_names(content: str) -> List[str]:
        """Extract party names from document content."""
        # Simple extraction - could be enhanced with NLP
        parties = []
        
        # Look for common patterns
        patterns = [
            r'(?:between|among)\s+([A-Z][a-zA-Z\s]+?)(?:\s+and\s+([A-Z][a-zA-Z\s]+?))?(?:\s*[\(,])',
            r'Landlord:\s*([A-Z][a-zA-Z\s]+)',
            r'Tenant:\s*([A-Z][a-zA-Z\s]+)',
            r'Employee:\s*([A-Z][a-zA-Z\s]+)',
            r'Employer:\s*([A-Z][a-zA-Z\s]+)',
            r'Client:\s*([A-Z][a-zA-Z\s]+)',
            r'Service Provider:\s*([A-Z][a-zA-Z\s]+)'
        ]
        
        for pattern in patterns:
            matches = re.findall(pattern, content, re.IGNORECASE)
            for match in matches:
                if isinstance(match, tuple):
                    parties.extend([m.strip() for m in match if m.strip()])
                else:
                    parties.append(match.strip())
        
        # Clean and deduplicate
        cleaned_parties = []
        for party in parties:
            if party and len(party) > 2 and party not in cleaned_parties:
                cleaned_parties.append(party)
        
        return cleaned_parties[:4]  # Limit to 4 parties max


class DocumentValidator:
    """Validates processed legal documents for completeness and quality."""
    
    def __init__(self):
        """Initialize document validator."""
        self.validation_rules = self._build_validation_rules()
    
    def _build_validation_rules(self) -> Dict[DocumentType, Dict[str, Any]]:
        """Build validation rules for different document types."""
        return {
            DocumentType.LEASE_AGREEMENT: {
                "required_sections": [
                    "parties", "property", "term", "rent", "deposit", "signature"
                ],
                "required_phrases": [
                    "landlord", "tenant", "rent", "lease", "property"
                ],
                "min_word_count": 400,
                "max_word_count": 3000,
                "must_have_signatures": True,
                "must_have_dates": True
            },
            
            DocumentType.NDA: {
                "required_sections": [
                    "parties", "confidential information", "obligations", "term", "signature"
                ],
                "required_phrases": [
                    "confidential", "disclosure", "information", "party"
                ],
                "min_word_count": 300,
                "max_word_count": 2000,
                "must_have_signatures": True,
                "must_have_dates": True
            },
            
            DocumentType.SERVICE_AGREEMENT: {
                "required_sections": [
                    "parties", "services", "payment", "term", "signature"
                ],
                "required_phrases": [
                    "service", "payment", "deliverable", "client"
                ],
                "min_word_count": 400,
                "max_word_count": 2500,
                "must_have_signatures": True,
                "must_have_dates": True
            },
            
            DocumentType.EMPLOYMENT_CONTRACT: {
                "required_sections": [
                    "parties", "position", "compensation", "term", "signature"
                ],
                "required_phrases": [
                    "employee", "employer", "position", "salary", "compensation"
                ],
                "min_word_count": 500,
                "max_word_count": 3500,
                "must_have_signatures": True,
                "must_have_dates": True
            }
        }
    
    def validate_document(self, content: str, document_type: DocumentType) -> Dict[str, Any]:
        """Validate document content against type-specific rules."""
        
        rules = self.validation_rules.get(document_type, {})
        results = {
            "is_valid": True,
            "score": 100,
            "errors": [],
            "warnings": [],
            "suggestions": [],
            "statistics": self._calculate_document_stats(content)
        }
        
        try:
            # Word count validation
            word_count = len(content.split())
            min_words = rules.get("min_word_count", 100)
            max_words = rules.get("max_word_count", 5000)
            
            if word_count < min_words:
                results["errors"].append(f"Document too short ({word_count} words, minimum {min_words})")
                results["score"] -= 20
                results["is_valid"] = False
            elif word_count > max_words:
                results["warnings"].append(f"Document very long ({word_count} words, maximum recommended {max_words})")
                results["score"] -= 5
            
            # Required phrases validation
            content_lower = content.lower()
            required_phrases = rules.get("required_phrases", [])
            missing_phrases = [phrase for phrase in required_phrases if phrase not in content_lower]
            
            if missing_phrases:
                results["warnings"].append(f"Missing key terms: {', '.join(missing_phrases)}")
                results["score"] -= len(missing_phrases) * 5
            
            # Section structure validation
            required_sections = rules.get("required_sections", [])
            missing_sections = []
            
            for section in required_sections:
                if not self._has_section(content, section):
                    missing_sections.append(section)
            
            if missing_sections:
                results["warnings"].append(f"Possibly missing sections: {', '.join(missing_sections)}")
                results["score"] -= len(missing_sections) * 10
            
            # Signature validation
            if rules.get("must_have_signatures", False):
                if not self._has_signature_block(content):
                    results["errors"].append("Missing signature block")
                    results["score"] -= 15
                    results["is_valid"] = False
            
            # Date validation
            if rules.get("must_have_dates", False):
                if not self._has_dates(content):
                    results["warnings"].append("Consider adding specific dates")
                    results["score"] -= 5
            
            # Legal disclaimer validation
            if "disclaimer" not in content_lower and "attorney" not in content_lower:
                results["suggestions"].append("Consider adding legal disclaimer")
                results["score"] -= 5
            
            # Placeholder validation
            placeholders = re.findall(r'\[PLACEHOLDER[^\]]*\]', content)
            if placeholders:
                results["warnings"].append(f"Document contains {len(placeholders)} placeholders that should be filled")
                results["score"] -= len(placeholders) * 3
            
            # Overall quality assessment
            results["quality_grade"] = self._calculate_quality_grade(results["score"])
            
        except Exception as e:
            logger.error(f"Document validation error: {str(e)}")
            results["errors"].append(f"Validation error: {str(e)}")
            results["is_valid"] = False
        
        return results
    
    def _calculate_document_stats(self, content: str) -> Dict[str, Any]:
        """Calculate document statistics."""
        words = content.split()
        sentences = re.split(r'[.!?]+', content)
        paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]
        
        return {
            "word_count": len(words),
            "sentence_count": len([s for s in sentences if s.strip()]),
            "paragraph_count": len(paragraphs),
            "character_count": len(content),
            "average_words_per_sentence": len(words) / max(len(sentences), 1),
            "reading_time_minutes": len(words) / 200  # Assume 200 WPM
        }
    
    def _has_section(self, content: str, section: str) -> bool:
        """Check if document has a specific section."""
        section_patterns = {
            "parties": r'(?i)\b(?:parties?|between|among)\b',
            "property": r'(?i)\b(?:property|premises|real estate|address)\b',
            "term": r'(?i)\b(?:term|duration|period|expires?)\b',
            "rent": r'(?i)\b(?:rent|rental|payment|monthly)\b',
            "deposit": r'(?i)\b(?:deposit|security)\b',
            "signature": r'(?i)\b(?:signature|signed|executed)\b',
            "confidential information": r'(?i)\b(?:confidential|proprietary|trade secret)\b',
            "obligations": r'(?i)\b(?:obligation|duty|responsibility|shall)\b',
            "services": r'(?i)\b(?:service|work|deliverable|perform)\b',
            "position": r'(?i)\b(?:position|role|job|title|employ)\b',
            "compensation": r'(?i)\b(?:compensation|salary|wage|pay)\b'
        }
        
        pattern = section_patterns.get(section.lower())
        if pattern:
            return bool(re.search(pattern, content))
        
        # Fallback: look for section as a header
        header_pattern = rf'(?i)^{re.escape(section)}:?\s*$'
        return bool(re.search(header_pattern, content, re.MULTILINE))
    
    def _has_signature_block(self, content: str) -> bool:
        """Check if document has signature block."""
        signature_patterns = [
            r'(?i)signature:?\s*_+',
            r'(?i)signed:?\s*_+',
            r'(?i)_+\s*signature',
            r'(?i)_+\s*date',
            r'(?i)by:\s*_+'
        ]
        
        return any(re.search(pattern, content) for pattern in signature_patterns)
    
    def _has_dates(self, content: str) -> bool:
        """Check if document contains date references."""
        date_patterns = [
            r'\d{1,2}/\d{1,2}/\d{4}',
            r'\d{4}-\d{2}-\d{2}',
            r'(?i)(?:january|february|march|april|may|june|july|august|september|october|november|december)\s+\d{1,2},?\s+\d{4}',
            r'(?i)date:?\s*_+',
            r'(?i)effective\s+date',
            r'(?i)expir\w+\s+date'
        ]
        
        return any(re.search(pattern, content) for pattern in date_patterns)
    
    def _calculate_quality_grade(self, score: int) -> str:
        """Calculate quality grade based on score."""
        if score >= 90:
            return "Excellent"
        elif score >= 80:
            return "Good"
        elif score >= 70:
            return "Fair"
        elif score >= 60:
            return "Needs Improvement"
        else:
            return "Poor"


class DocumentFormatter:
    """Handles document formatting for different output formats."""
    
    def __init__(self):
        """Initialize document formatter."""
        self.temp_dir = Path(tempfile.gettempdir()) / "legal_doc_generator"
        self.temp_dir.mkdir(exist_ok=True)
    
    def format_document(self, content: str, format: DocumentFormat, 
                       metadata: Dict[str, Any] = None) -> ProcessedDocument:
        """Format document to specified output format."""
        
        start_time = datetime.now()
        warnings = []
        
        try:
            if format == DocumentFormat.PDF:
                if not REPORTLAB_AVAILABLE:
                    raise FormatNotSupportedError("PDF generation requires ReportLab library")
                return self._generate_pdf(content, metadata, warnings)
            
            elif format == DocumentFormat.DOCX:
                if not PYTHON_DOCX_AVAILABLE:
                    raise FormatNotSupportedError("DOCX generation requires python-docx library")
                return self._generate_docx(content, metadata, warnings)
            
            elif format == DocumentFormat.HTML:
                return self._generate_html(content, metadata, warnings)
            
            elif format == DocumentFormat.TXT:
                return self._generate_txt(content, metadata, warnings)
            
            elif format == DocumentFormat.RTF:
                return self._generate_rtf(content, metadata, warnings)
            
            else:
                raise FormatNotSupportedError(f"Format {format.value} not supported")
        
        except Exception as e:
            logger.error(f"Document formatting error: {str(e)}")
            raise PostprocessingError(f"Failed to format document: {str(e)}")
        
        finally:
            processing_time = (datetime.now() - start_time).total_seconds()
    
    def _generate_pdf(self, content: str, metadata: Dict[str, Any] = None, 
                     warnings: List[str] = None) -> ProcessedDocument:
        """Generate PDF document using ReportLab."""
        
        # Create temporary file
        pdf_path = self.temp_dir / f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        # Create PDF document
        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=letter,
            rightMargin=72, leftMargin=72,
            topMargin=72, bottomMargin=72
        )
        
        # Get styles
        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            alignment=TA_CENTER,
            spaceAfter=30,
            fontSize=16,
            bold=True
        )
        
        body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['Normal'],
            alignment=TA_JUSTIFY,
            spaceAfter=12,
            fontSize=10,
            leading=14
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            spaceAfter=12,
            fontSize=12,
            bold=True
        )
        
        # Build story (content)
        story = []
        
        # Add title if available
        title = metadata.get('title', 'Legal Document') if metadata else 'Legal Document'
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 20))
        
        # Process content
        paragraphs = content.split('\n\n')
        
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            
            # Check if it's a heading (ALL CAPS or starts with number)
            if para.isupper() and len(para) < 100:
                story.append(Paragraph(para, heading_style))
            elif re.match(r'^\d+\.', para):
                story.append(Paragraph(para, heading_style))
            else:
                # Regular paragraph
                # Escape HTML characters and preserve formatting
                para = para.replace('<', '&lt;').replace('>', '&gt;')
                para = para.replace('\n', '<br/>')
                story.append(Paragraph(para, body_style))
            
            story.append(Spacer(1, 6))
        
        # Build PDF
        doc.build(story)
        
        # Get file size
        file_size = pdf_path.stat().st_size
        
        return ProcessedDocument(
            content=content,
            format=DocumentFormat.PDF,
            file_path=str(pdf_path),
            file_size=file_size,
            metadata=metadata,
            warnings=warnings or []
        )
    
    def _generate_docx(self, content: str, metadata: Dict[str, Any] = None, 
                      warnings: List[str] = None) -> ProcessedDocument:
        """Generate DOCX document using python-docx."""
        
        # Create temporary file
        docx_path = self.temp_dir / f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        
        # Create document
        doc = DocxDocument()
        
        # Add title
        title = metadata.get('title', 'Legal Document') if metadata else 'Legal Document'
        title_para = doc.add_paragraph(title)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.runs[0]
        title_run.font.size = Inches(0.2)
        title_run.bold = True
        
        # Add content
        paragraphs = content.split('\n\n')
        
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            
            # Add paragraph
            doc_para = doc.add_paragraph(para)
            
            # Check if it's a heading
            if para.isupper() and len(para) < 100:
                doc_para.runs[0].bold = True
                doc_para.runs[0].font.size = Inches(0.15)
            elif re.match(r'^\d+\.', para):
                doc_para.runs[0].bold = True
        
        # Save document
        doc.save(str(docx_path))
        
        # Get file size
        file_size = docx_path.stat().st_size
        
        return ProcessedDocument(
            content=content,
            format=DocumentFormat.DOCX,
            file_path=str(docx_path),
            file_size=file_size,
            metadata=metadata,
            warnings=warnings or []
        )
    
    def _generate_html(self, content: str, metadata: Dict[str, Any] = None, 
                      warnings: List[str] = None) -> ProcessedDocument:
        """Generate HTML document."""
        
        title = metadata.get('title', 'Legal Document') if metadata else 'Legal Document'
        
        # Convert content to HTML
        html_content = self._text_to_html(content)
        
        html_template = f"""
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        body {{
            font-family: 'Times New Roman', Times, serif;
            line-height: 1.6;
            max-width: 800px;
            margin: 0 auto;
            padding: 20px;
            color: #333;
        }}
        h1 {{
            text-align: center;
            color: #2c3e50;
            border-bottom: 2px solid #3498db;
            padding-bottom: 10px;
        }}
        h2 {{
            color: #2c3e50;
            margin-top: 30px;
        }}
        p {{
            text-align: justify;
            margin-bottom: 15px;
        }}
        .signature-block {{
            margin-top: 40px;
            border-top: 1px solid #ddd;
            padding-top: 20px;
        }}
        .disclaimer {{
            background-color: #f8f9fa;
            border: 1px solid #dee2e6;
            padding: 15px;
            margin: 20px 0;
            border-radius: 5px;
            font-size: 0.9em;
        }}
        @media print {{
            body {{ margin: 0; }}
            .no-print {{ display: none; }}
        }}
    </style>
</head>
<body>
    <h1>{title}</h1>
    {html_content}
    <div class="no-print" style="text-align: center; margin-top: 30px;">
        <button onclick="window.print()">Print Document</button>
    </div>
</body>
</html>
"""
        
        # Create temporary file
        html_path = self.temp_dir / f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
        
        with open(html_path, 'w', encoding='utf-8') as f:
            f.write(html_template)
        
        file_size = html_path.stat().st_size
        
        return ProcessedDocument(
            content=html_template,
            format=DocumentFormat.HTML,
            file_path=str(html_path),
            file_size=file_size,
            metadata=metadata,
            warnings=warnings or []
        )
    
    def _generate_txt(self, content: str, metadata: Dict[str, Any] = None, 
                     warnings: List[str] = None) -> ProcessedDocument:
        """Generate plain text document."""
        
        title = metadata.get('title', 'Legal Document') if metadata else 'Legal Document'
        
        # Format as plain text with proper structure
        txt_content = f"{title.upper()}\n{'=' * len(title)}\n\n{content}"
        
        # Create temporary file
        txt_path = self.temp_dir / f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
        
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write(txt_content)
        
        file_size = txt_path.stat().st_size
        
        return ProcessedDocument(
            content=txt_content,
            format=DocumentFormat.TXT,
            file_path=str(txt_path),
            file_size=file_size,
            metadata=metadata,
            warnings=warnings or []
        )
    
    def _generate_rtf(self, content: str, metadata: Dict[str, Any] = None, 
                     warnings: List[str] = None) -> ProcessedDocument:
        """Generate RTF document."""
        
        title = metadata.get('title', 'Legal Document') if metadata else 'Legal Document'
        
        # Basic RTF template
        rtf_content = r"""{\rtf1\ansi\deff0 {\fonttbl {\f0 Times New Roman;}}
{\colortbl;\red0\green0\blue0;}
\f0\fs24
""" + f"""
\\pard\\qc\\b\\fs28 {title}\\b0\\fs24\\par
\\par
\\pard\\qj
"""
        
        # Convert content to RTF
        rtf_text = content.replace('\n\n', '\\par\\par ')
        rtf_text = rtf_text.replace('\n', ' ')
        rtf_text = rtf_text.replace('\\', '\\\\')
        rtf_text = rtf_text.replace('{', '\\{')
        rtf_text = rtf_text.replace('}', '\\}')
        
        rtf_content += rtf_text + "\\par}"
        
        # Create temporary file
        rtf_path = self.temp_dir / f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.rtf"
        
        with open(rtf_path, 'w', encoding='utf-8') as f:
            f.write(rtf_content)
        
        file_size = rtf_path.stat().st_size
        
        return ProcessedDocument(
            content=rtf_content,
            format=DocumentFormat.RTF,
            file_path=str(rtf_path),
            file_size=file_size,
            metadata=metadata,
            warnings=warnings or []
        )
    
    def _text_to_html(self, text: str) -> str:
        """Convert plain text to HTML with basic formatting."""
        
        # Escape HTML characters
        text = text.replace('&', '&amp;')
        text = text.replace('<', '&lt;')
        text = text.replace('>', '&gt;')
        text = text.replace('"', '&quot;')
        
        # Convert paragraphs
        paragraphs = text.split('\n\n')
        html_parts = []
        
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            
            # Check if it's a heading
            if para.isupper() and len(para) < 100:
                html_parts.append(f'<h2>{para}</h2>')
            elif re.match(r'^\d+\.', para):
                html_parts.append(f'<h3>{para}</h3>')
            elif 'DISCLAIMER' in para.upper():
                html_parts.append(f'<div class="disclaimer">{para.replace(chr(10), "<br>")}</div>')
            elif 'Signature:' in para or '_____' in para:
                html_parts.append(f'<div class="signature-block">{para.replace(chr(10), "<br>")}</div>')
            else:
                # Regular paragraph
                html_parts.append(f'<p>{para.replace(chr(10), "<br>")}</p>')
        
        return '\n'.join(html_parts)
    
    def cleanup_temp_files(self, older_than_hours: int = 24):
        """Clean up temporary files older than specified hours."""
        try:
            import time
            current_time = time.time()
            cutoff_time = current_time - (older_than_hours * 3600)
            
            for file_path in self.temp_dir.glob('*'):
                if file_path.is_file() and file_path.stat().st_mtime < cutoff_time:
                    file_path.unlink()
                    logger.debug(f"Cleaned up temporary file: {file_path}")
        
        except Exception as e:
            logger.error(f"Error cleaning up temporary files: {str(e)}")


class PostprocessingService:
    """Main postprocessing service that orchestrates all postprocessing steps."""
    
    def __init__(self):
        """Initialize postprocessing components."""
        self.cleaner = DocumentCleaner()
        self.validator = DocumentValidator()
        self.formatter = DocumentFormatter()
    
    def process_document(self, raw_content: str, document_type: DocumentType,
                        output_format: DocumentFormat, metadata: Dict[str, Any] = None,
                        include_disclaimer: bool = True, 
                        include_signatures: bool = True) -> ProcessedDocument:
        """
        Complete postprocessing pipeline for generated documents.
        
        Args:
            raw_content: Raw content from LLM generation
            document_type: Type of legal document
            output_format: Desired output format
            metadata: Additional document metadata
            include_disclaimer: Whether to add legal disclaimer
            include_signatures: Whether to add signature blocks
            
        Returns:
            ProcessedDocument: Fully processed document
        """
        
        logger.info(f"Starting postprocessing for {document_type.value} -> {output_format.value}")
        start_time = datetime.now()
        
        try:
            # Step 1: Clean the content
            cleaned_content = self.cleaner.clean_document_content(raw_content)
            
            # Step 2: Add legal disclaimer if requested
            if include_disclaimer:
                cleaned_content = self.cleaner.add_legal_disclaimer(cleaned_content, document_type)
            
            # Step 3: Add signature blocks if requested
            if include_signatures:
                parties = metadata.get('parties', []) if metadata else []
                cleaned_content = self.cleaner.format_signature_blocks(cleaned_content, parties)
            
            # Step 4: Validate the document
            validation_results = self.validator.validate_document(cleaned_content, document_type)
            
            # Step 5: Format to desired output
            processed_doc = self.formatter.format_document(
                cleaned_content, 
                output_format, 
                metadata
            )
            
            # Step 6: Add validation results and processing time
            processing_time = (datetime.now() - start_time).total_seconds()
            processed_doc.validation_results = validation_results
            processed_doc.processing_time = processing_time
            
            # Add warnings from validation
            if validation_results.get('warnings'):
                processed_doc.warnings.extend(validation_results['warnings'])
            
            # Add processing metadata
            if not processed_doc.metadata:
                processed_doc.metadata = {}
            
            processed_doc.metadata.update({
                'processing_timestamp': datetime.now(timezone.utc).isoformat(),
                'document_type': document_type.value,
                'processing_time_seconds': processing_time,
                'validation_score': validation_results.get('score', 0),
                'quality_grade': validation_results.get('quality_grade', 'Unknown'),
                'word_count': validation_results.get('statistics', {}).get('word_count', 0),
                'character_count': validation_results.get('statistics', {}).get('character_count', 0)
            })
            
            logger.info(f"Postprocessing completed in {processing_time:.2f}s, "
                       f"validation score: {validation_results.get('score', 0)}/100")
            
            return processed_doc
            
        except Exception as e:
            logger.error(f"Postprocessing failed: {str(e)}")
            raise PostprocessingError(f"Document postprocessing failed: {str(e)}")
    
    def validate_only(self, content: str, document_type: DocumentType) -> Dict[str, Any]:
        """Validate document content without full processing."""
        return self.validator.validate_document(content, document_type)
    
    def format_only(self, content: str, output_format: DocumentFormat, 
                   metadata: Dict[str, Any] = None) -> ProcessedDocument:
        """Format document content without cleaning or validation."""
        return self.formatter.format_document(content, output_format, metadata)
    
    def get_supported_formats(self) -> List[DocumentFormat]:
        """Get list of supported output formats."""
        supported = [DocumentFormat.TXT, DocumentFormat.HTML, DocumentFormat.RTF]
        
        if REPORTLAB_AVAILABLE:
            supported.append(DocumentFormat.PDF)
        
        if PYTHON_DOCX_AVAILABLE:
            supported.append(DocumentFormat.DOCX)
        
        return supported
    
    def cleanup_old_files(self):
        """Clean up old temporary files."""
        self.formatter.cleanup_temp_files()
    
    def get_processing_stats(self) -> Dict[str, Any]:
        """Get postprocessing service statistics."""
        return {
            'supported_formats': [fmt.value for fmt in self.get_supported_formats()],
            'supported_document_types': [dt.value for dt in DocumentType],
            'libraries': {
                'reportlab_available': REPORTLAB_AVAILABLE,
                'python_docx_available': PYTHON_DOCX_AVAILABLE
            },
            'temp_directory': str(self.formatter.temp_dir),
            'validation_rules_count': len(self.validator.validation_rules)
        }


# Global service instance
postprocessing_service = PostprocessingService()


# ==================== UTILITY FUNCTIONS ====================

def process_generated_document(content: str, document_type: DocumentType,
                             output_format: DocumentFormat = DocumentFormat.PDF,
                             metadata: Dict[str, Any] = None) -> ProcessedDocument:
    """
    High-level function for document postprocessing.
    
    Args:
        content: Generated document content
        document_type: Type of legal document
        output_format: Desired output format
        metadata: Document metadata
        
    Returns:
        ProcessedDocument: Processed document
    """
    return postprocessing_service.process_document(
        content, document_type, output_format, metadata
    )


def validate_document_quality(content: str, document_type: DocumentType) -> Dict[str, Any]:
    """
    Quick validation of document quality.
    
    Args:
        content: Document content to validate
        document_type: Type of document
        
    Returns:
        Dict: Validation results
    """
    return postprocessing_service.validate_only(content, document_type)


def convert_document_format(content: str, from_format: DocumentFormat, 
                          to_format: DocumentFormat, metadata: Dict[str, Any] = None) -> ProcessedDocument:
    """
    Convert document from one format to another.
    
    Args:
        content: Document content
        from_format: Source format (currently only supports text input)
        to_format: Target format
        metadata: Document metadata
        
    Returns:
        ProcessedDocument: Converted document
    """
    if from_format != DocumentFormat.TXT:
        raise FormatNotSupportedError("Currently only text to other format conversion is supported")
    
    return postprocessing_service.format_only(content, to_format, metadata)


def calculate_document_hash(content: str) -> str:
    """
    Calculate SHA-256 hash of document content for integrity checking.
    
    Args:
        content: Document content
        
    Returns:
        str: SHA-256 hash in hexadecimal format
    """
    return hashlib.sha256(content.encode('utf-8')).hexdigest()


def estimate_processing_time(content_length: int, output_format: DocumentFormat) -> float:
    """
    Estimate processing time based on content length and format.
    
    Args:
        content_length: Length of content in characters
        output_format: Target output format
        
    Returns:
        float: Estimated processing time in seconds
    """
    # Base processing time (cleaning, validation)
    base_time = 0.1 + (content_length / 10000) * 0.5
    
    # Format-specific overhead
    format_overhead = {
        DocumentFormat.TXT: 0.1,
        DocumentFormat.HTML: 0.2,
        DocumentFormat.RTF: 0.3,
        DocumentFormat.PDF: 1.0,  # PDF generation is slower
        DocumentFormat.DOCX: 0.8
    }
    
    overhead = format_overhead.get(output_format, 0.5)
    
    return base_time + overhead


def get_format_mime_type(format: DocumentFormat) -> str:
    """
    Get MIME type for document format.
    
    Args:
        format: Document format
        
    Returns:
        str: MIME type
    """
    mime_types = {
        DocumentFormat.PDF: 'application/pdf',
        DocumentFormat.DOCX: 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        DocumentFormat.HTML: 'text/html',
        DocumentFormat.TXT: 'text/plain',
        DocumentFormat.RTF: 'application/rtf'
    }
    
    return mime_types.get(format, 'application/octet-stream')


def get_format_file_extension(format: DocumentFormat) -> str:
    """
    Get file extension for document format.
    
    Args:
        format: Document format
        
    Returns:
        str: File extension with dot
    """
    extensions = {
        DocumentFormat.PDF: '.pdf',
        DocumentFormat.DOCX: '.docx',
        DocumentFormat.HTML: '.html',
        DocumentFormat.TXT: '.txt',
        DocumentFormat.RTF: '.rtf'
    }
    
    return extensions.get(format, '.txt')


# ==================== BATCH PROCESSING ====================

class BatchProcessor:
    """Handle batch processing of multiple documents."""
    
    def __init__(self, max_concurrent: int = 3):
        """Initialize batch processor."""
        self.max_concurrent = max_concurrent
        self.service = postprocessing_service
    
    async def process_batch(self, documents: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        Process multiple documents concurrently.
        
        Args:
            documents: List of document processing requests
            
        Returns:
            List of processing results
        """
        import asyncio
        
        # Limit concurrency to prevent resource exhaustion
        semaphore = asyncio.Semaphore(self.max_concurrent)
        
        async def process_single(doc_request):
            async with semaphore:
                try:
                    result = await asyncio.to_thread(
                        self.service.process_document,
                        doc_request['content'],
                        DocumentType(doc_request['document_type']),
                        DocumentFormat(doc_request['output_format']),
                        doc_request.get('metadata')
                    )
                    
                    return {
                        'success': True,
                        'document_id': doc_request.get('id'),
                        'result': result,
                        'processing_time': result.processing_time
                    }
                    
                except Exception as e:
                    return {
                        'success': False,
                        'document_id': doc_request.get('id'),
                        'error': str(e),
                        'processing_time': 0
                    }
        
        # Process all documents concurrently
        tasks = [process_single(doc) for doc in documents]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        
        return results


# Global batch processor instance
batch_processor = BatchProcessor()


# ==================== TESTING UTILITIES ====================

def test_postprocessing(sample_content: str = None) -> Dict[str, Any]:
    """
    Test postprocessing functionality with sample content.
    
    Args:
        sample_content: Optional sample content to test
        
    Returns:
        Dict: Test results
    """
    if not sample_content:
        sample_content = """
RESIDENTIAL LEASE AGREEMENT

This lease agreement is entered into between John Smith (Landlord) and Jane Doe (Tenant) 
for the property located at 123 Main Street, Anytown, State 12345.

TERM: The lease term is 12 months beginning January 1, 2024 and ending December 31, 2024.

RENT: The monthly rent is $1,500.00 due on the first day of each month.

SECURITY DEPOSIT: Tenant shall pay a security deposit of $3,000.00.

The parties agree to the terms set forth in this agreement.

Signature: _________________________  Date: ________________
Landlord: John Smith

Signature: _________________________  Date: ________________
Tenant: Jane Doe
"""
    
    test_results = {}
    
    try:
        # Test validation
        validation = postprocessing_service.validate_only(sample_content, DocumentType.LEASE_AGREEMENT)
        test_results['validation'] = {
            'score': validation.get('score', 0),
            'is_valid': validation.get('is_valid', False),
            'errors': len(validation.get('errors', [])),
            'warnings': len(validation.get('warnings', []))
        }
        
        # Test format support
        supported_formats = postprocessing_service.get_supported_formats()
        test_results['supported_formats'] = [f.value for f in supported_formats]
        
        # Test processing for each supported format
        format_results = {}
        for fmt in supported_formats:
            try:
                start_time = datetime.now()
                result = postprocessing_service.process_document(
                    sample_content,
                    DocumentType.LEASE_AGREEMENT,
                    fmt,
                    {'title': 'Test Lease Agreement'}
                )
                processing_time = (datetime.now() - start_time).total_seconds()
                
                format_results[fmt.value] = {
                    'success': True,
                    'processing_time': processing_time,
                    'file_size': result.file_size,
                    'has_file': result.file_path is not None,
                    'warnings': len(result.warnings or [])
                }
                
                # Clean up test file
                if result.file_path and os.path.exists(result.file_path):
                    os.unlink(result.file_path)
                    
            except Exception as e:
                format_results[fmt.value] = {
                    'success': False,
                    'error': str(e)
                }
        
        test_results['format_processing'] = format_results
        test_results['overall_success'] = True
        
    except Exception as e:
        test_results['overall_success'] = False
        test_results['error'] = str(e)
    
    return test_results